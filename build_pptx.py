import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors
    c_bg_dark = RGBColor(15, 23, 42)       # Slate 900
    c_card_bg = RGBColor(30, 41, 59)       # Slate 800
    c_primary = RGBColor(37, 99, 235)      # Blue 600
    c_accent = RGBColor(14, 165, 233)      # Sky 500
    c_accent_green = RGBColor(16, 185, 129)# Emerald 500
    c_text_white = RGBColor(255, 255, 255)
    c_text_muted = RGBColor(148, 163, 184) # Slate 400
    c_border = RGBColor(51, 65, 85)        # Slate 700

    def add_bg(slide):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        shape.fill.solid()
        shape.fill.fore_color.rgb = c_bg_dark
        shape.line.fill.background()
        return shape

    def add_header(slide, title_text, category="AUTOMATIZACIÓN DE AGENTES DE IA | IDAT - PACHAQTEC"):
        # Category label (larger font)
        tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.7), Inches(0.4))
        p_cat = tb_cat.text_frame.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(13)
        p_cat.font.bold = True
        p_cat.font.color.rgb = c_accent

        # Title (larger font 28pt)
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.68), Inches(11.7), Inches(0.85))
        p_title = tb_title.text_frame.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = c_text_white

        # Divider line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.55), Inches(11.733), Inches(0.025))
        line.fill.solid()
        line.fill.fore_color.rgb = c_border
        line.line.fill.background()

    def add_speaker_note(slide, note_text):
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = note_text

    # =========================================================================
    # SLIDE 1: PORTADA
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1)

    # Accent decorative badge
    dec = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(5.2), Inches(0.55))
    dec.fill.solid()
    dec.fill.fore_color.rgb = c_primary
    dec.line.fill.background()
    p_dec = dec.text_frame.paragraphs[0]
    p_dec.text = "CURSO: AUTOMATIZACIÓN DE AGENTES DE IA"
    p_dec.font.size = Pt(14)
    p_dec.font.bold = True
    p_dec.font.color.rgb = c_text_white
    p_dec.alignment = PP_ALIGN.CENTER

    tb1 = s1.shapes.add_textbox(Inches(0.8), Inches(2.1), Inches(11.7), Inches(3.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "Casos Reales y Práctica con OpenAI API"
    p1.font.size = Pt(46)
    p1.font.bold = True
    p1.font.color.rgb = c_text_white

    p1_sub = tf1.add_paragraph()
    p1_sub.text = "Subindicador 3: Aplicación práctica de agentes en atención al cliente, generación de contenido, automatización de procesos y productividad."
    p1_sub.font.size = Pt(22)
    p1_sub.font.color.rgb = c_text_muted
    p1_sub.space_before = Pt(12)

    # Meta card
    meta = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.5), Inches(11.7), Inches(1.4))
    meta.fill.solid()
    meta.fill.fore_color.rgb = c_card_bg
    meta.line.color.rgb = c_border
    tf_meta = meta.text_frame
    tf_meta.word_wrap = True
    p_m1 = tf_meta.paragraphs[0]
    p_m1.text = "Docente: IDAT / Pachaqtec   •   Horario: 19:30 - 21:45 hrs (135 min)"
    p_m1.font.size = Pt(17)
    p_m1.font.bold = True
    p_m1.font.color.rgb = c_accent
    p_m2 = tf_meta.add_paragraph()
    p_m2.text = "Entorno Hands-On: OpenAI Platform (Playground & Chat Completions API)"
    p_m2.font.size = Pt(15)
    p_m2.font.color.rgb = c_text_white
    p_m2.space_before = Pt(4)

    add_speaker_note(s1, "Iniciar a las 19:30 en punto. Dar la bienvenida y recordar que hoy es la sesión práctica donde aterrizaremos la teoría de los agentes a la plataforma real de OpenAI.")

    # =========================================================================
    # SLIDE 2: AGENDA DE HOY
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2)
    add_header(s2, "Agenda de la Sesión de Hoy (19:30 - 21:45)")

    agenda_items = [
        ("19:30 - 19:42", "Apertura y Conexión", "Recap de subindicadores 1 y 2: Tipos de agentes y arquitectura básica."),
        ("19:42 - 20:10", "Los 4 Casos de Uso Clave", "Atención al cliente, generación de contenido, procesos (JSON) y productividad."),
        ("20:10 - 20:25", "OpenAI Platform & Playground", "ChatGPT vs API, System Prompt, Temperature, Tokens y Modelos."),
        ("20:25 - 20:30", "Pausa Activa / Acceso Técnico", "Verificación de cuentas en platform.openai.com para todos los alumnos."),
        ("20:30 - 21:20", "Hands-On: 4 Prácticas en Vivo", "Configuración y prueba de agentes reales paso a paso en el Playground."),
        ("21:20 - 21:45", "Reto Aplicado y Cierre", "Construcción de un agente propio y conexión con Subindicador 4.")
    ]

    for idx, (hor, tit, desc) in enumerate(agenda_items):
        col = idx % 2
        row = idx // 2
        x = Inches(0.8 + col * 5.95)
        y = Inches(1.75 + row * 1.75)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.75), Inches(1.55))
        card.fill.solid()
        card.fill.fore_color.rgb = c_card_bg
        card.line.color.rgb = c_border
        tf = card.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = f"{hor}  •  {tit}"
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = c_accent
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(14)
        p_d.font.color.rgb = c_text_white
        p_d.space_before = Pt(4)

    add_speaker_note(s2, "Presentar los bloques de tiempo para que los alumnos sepan cuándo será la parte teórica y cuándo empezarán a construir prompts en vivo.")

    # =========================================================================
    # SLIDE 3: RECAP SUBINDICADORES 1 Y 2
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3)
    add_header(s3, "Conexión: ¿Qué aprendimos el martes? (Subindicadores 1 y 2)")

    recap_cols = [
        ("Subindicador 1: Tipos de Agentes", [
            "• Reactivo: Responde a estímulos inmediatos sin memoria histórica.",
            "• Deliberativo: Posee un modelo del mundo, planifica y evalúa opciones antes de actuar.",
            "• Autónomo: Toma decisiones, usa herramientas y persigue metas sin supervisión continua."
        ]),
        ("Subindicador 2: Arquitectura del Agente", [
            "• Input (Entrada): Texto, audios, documentos o eventos de webhook.",
            "• LLM + Lógica: Procesa el contexto guiado por el System Prompt.",
            "• Output (Acción): Respuesta en texto, llamada a API o JSON estructurado para otro software."
        ])
    ]

    for i, (col_title, points) in enumerate(recap_cols):
        x = Inches(0.8 + i * 5.95)
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.75), Inches(5.75), Inches(5.1))
        card.fill.solid()
        card.fill.fore_color.rgb = c_card_bg
        card.line.color.rgb = c_border
        tf = card.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = col_title
        p_t.font.size = Pt(20)
        p_t.font.bold = True
        p_t.font.color.rgb = c_primary

        for pt in points:
            p = tf.add_paragraph()
            p.text = pt
            p.font.size = Pt(16)
            p.font.color.rgb = c_text_white
            p.space_before = Pt(16)

    add_speaker_note(s3, "Recordar a los alumnos que el LLM es el motor, pero el System Prompt es la personalidad y manual de operaciones del agente.")

    # =========================================================================
    # SLIDE 4: SUBINDICADOR 3 - PANORAMA GENERAL
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4)
    add_header(s4, "Subindicador 3: Los 4 Ejes de Aplicación Real en Organizaciones")

    ejes = [
        ("1. Atención al Cliente", "Resolución de consultas frecuentes, contención emocional, seguimiento de pedidos y escalamiento controlado a agentes humanos.", c_primary),
        ("2. Generación de Contenido", "Adaptación de mensajes a múltiples canales (LinkedIn, WhatsApp, Email) manteniendo la voz de marca y consistencia.", c_accent),
        ("3. Automatización de Procesos", "Extracción de entidades desde texto no estructurado hacia JSON para alimentar bases de datos (Airtable, Supabase, ERPs).", c_accent_green),
        ("4. Productividad y Análisis", "Conversión de notas dispersas o transcripciones en minutas ejecutivas, tareas accionables y diagramación inicial.", c_primary)
    ]

    for idx, (title, desc, color) in enumerate(ejes):
        x = Inches(0.8 + (idx % 2) * 5.95)
        y = Inches(1.75 + (idx // 2) * 2.65)
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.75), Inches(2.45))
        card.fill.solid()
        card.fill.fore_color.rgb = c_card_bg
        card.line.color.rgb = color
        card.line.width = Pt(2.5)
        tf = card.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(20)
        p_t.font.bold = True
        p_t.font.color.rgb = color
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(15)
        p_d.font.color.rgb = c_text_white
        p_d.space_before = Pt(8)

    add_speaker_note(s4, "Hacer énfasis en que cada uno de estos 4 ejes tendrá su ejercicio práctico en el Playground durante la segunda mitad de la clase.")

    # =========================================================================
    # SLIDE 5: CASO 1 - ATENCIÓN AL CLIENTE
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5)
    add_header(s5, "Eje 1: Automatización de Atención al Cliente (Tier 1 Support)")

    c1_left = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.75), Inches(5.75), Inches(5.1))
    c1_left.fill.solid()
    c1_left.fill.fore_color.rgb = c_card_bg
    c1_left.line.color.rgb = c_border
    tf1 = c1_left.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "¿Por qué implementar este Agente?"
    p1.font.size = Pt(19)
    p1.font.bold = True
    p1.font.color.rgb = c_accent

    bullet_pts = [
        "• Reducción del tiempo de respuesta (FRT) a < 2 segundos.",
        "• Cobertura 24/7 sin sobrecosto de turnos nocturnos.",
        "• Resolución inmediata del 60-70% de consultas repetitivas.",
        "• El equipo humano se enfoca en casos de alto valor o complejidad."
    ]
    for b in bullet_pts:
        p = tf1.add_paragraph()
        p.text = b
        p.font.size = Pt(15)
        p.font.color.rgb = c_text_white
        p.space_before = Pt(12)

    c1_right = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.75), Inches(1.75), Inches(5.75), Inches(5.1))
    c1_right.fill.solid()
    c1_right.fill.fore_color.rgb = c_card_bg
    c1_right.line.color.rgb = c_border
    tf2 = c1_right.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = "Reglas de Diseño (Guardrails):"
    p2.font.size = Pt(19)
    p2.font.bold = True
    p2.font.color.rgb = c_accent_green

    guardrails = [
        "1. Políticas explícitas: Plazos de devolución, costos de envío y exclusiones claras.",
        "2. Manejo de objeciones: Empatía validada sin prometer excepciones prohibidas.",
        "3. Protocolo de escalamiento: Código trigger para transferir al humano (ej. [DERIVAR_A_HUMANO]).",
        "4. Temperatura baja (0.1 - 0.2): Cero invención de datos o precios."
    ]
    for g in guardrails:
        p = tf2.add_paragraph()
        p.text = g
        p.font.size = Pt(15)
        p.font.color.rgb = c_text_white
        p.space_before = Pt(12)

    add_speaker_note(s5, "Preguntar a los alumnos: ¿Quién en su empresa tiene que responder lo mismo 20 veces al día por WhatsApp o correo?")

    # =========================================================================
    # SLIDE 6: CASO 2 - GENERACIÓN DE CONTENIDO
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6)
    add_header(s6, "Eje 2: Generación y Adaptación de Contenido Multicanal")

    items_c2 = [
        ("El reto en la empresa", "Un mismo mensaje o lanzamiento comercial debe comunicarse en LinkedIn, Blog, Email masivo y WhatsApp. Hacerlo manualmente lleva horas y suele perder el tono adecuado en cada red."),
        ("La solución del Agente", "Un agente con System Prompt que actúa como director editorial: recibe 1 solo hecho o hito y devuelve en 1 segundo las versiones formateadas para cada canal."),
        ("Control de Tono y Estilo", "Se definen reglas precisas: LinkedIn usa párrafos breves y llamada a la reflexión; WhatsApp usa viñetas con emojis y enlace directo; Email usa estructura de persuasión AIDA.")
    ]

    for idx, (t, d) in enumerate(items_c2):
        y = Inches(1.75 + idx * 1.75)
        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.7), Inches(1.55))
        card.fill.solid()
        card.fill.fore_color.rgb = c_card_bg
        card.line.color.rgb = c_border
        tf = card.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = t
        p_t.font.size = Pt(18)
        p_t.font.bold = True
        p_t.font.color.rgb = c_accent
        p_d = tf.add_paragraph()
        p_d.text = d
        p_d.font.size = Pt(15)
        p_d.font.color.rgb = c_text_white
        p_d.space_before = Pt(4)

    add_speaker_note(s6, "Explicar que aquí la temperatura puede ser un poco más alta (0.7) para permitir fluidez verbal y creatividad.")

    # =========================================================================
    # SLIDE 7: CASO 3 - PROCESOS (EL PODER DEL JSON)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7)
    add_header(s7, "Eje 3: Automatización de Procesos (Extracción y Estructuración JSON)")

    left = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.75), Inches(5.75), Inches(5.1))
    left.fill.solid()
    left.fill.fore_color.rgb = c_card_bg
    left.line.color.rgb = c_border
    tfl = left.text_frame
    tfl.word_wrap = True
    pl = tfl.paragraphs[0]
    pl.text = "Del texto libre a la Base de Datos"
    pl.font.size = Pt(19)
    pl.font.bold = True
    pl.font.color.rgb = c_accent

    txts = [
        "• Los humanos nos comunicamos en texto desordenado (un correo de reclamo, un audio, una cotización).",
        "• El software empresarial (ERPs, Make, Supabase, Airtable) necesita datos estructurados con nombres de columnas fijos.",
        "• El LLM actúa como traductor inteligente: lee lenguaje humano y devuelve un JSON válido.",
        "• ¡Esto elimina por completo el tipeo manual!"
    ]
    for t in txts:
        p = tfl.add_paragraph()
        p.text = t
        p.font.size = Pt(15)
        p.font.color.rgb = c_text_white
        p.space_before = Pt(12)

    # Right box: JSON example
    right = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.75), Inches(1.75), Inches(5.75), Inches(5.1))
    right.fill.solid()
    right.fill.fore_color.rgb = RGBColor(10, 15, 30)
    right.line.color.rgb = c_accent_green
    tfr = right.text_frame
    tfr.word_wrap = True
    pr = tfr.paragraphs[0]
    pr.text = "Salida del Agente (JSON Estructurado):"
    pr.font.size = Pt(16)
    pr.font.bold = True
    pr.font.color.rgb = c_accent_green

    code_lines = [
        '{',
        '  "cliente": "Carlos Mendoza",',
        '  "pedido": "PE-89412",',
        '  "categoria": "PAGOS",',
        '  "monto": 450.00,',
        '  "urgencia": "CRITICA",',
        '  "sentimiento": "MUY_MOLESTO",',
        '  "accion": "Devolver duplicado"',
        '}'
    ]
    for cl in code_lines:
        p = tfr.add_paragraph()
        p.text = cl
        p.font.size = Pt(14)
        p.font.name = "Courier New"
        p.font.color.rgb = RGBColor(226, 232, 240)
        p.space_before = Pt(2)

    add_speaker_note(s7, "Destacar que esta es la habilidad más codiciada para automatizadores no-code: convertir texto en variables para Make o Zapier.")

    # =========================================================================
    # SLIDE 8: CASO 4 - PRODUCTIVIDAD Y ANÁLISIS
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_bg(s8)
    add_header(s8, "Eje 4: Creación de Aplicaciones para Productividad y Gestión")

    cards_p4 = [
        ("De Notas Sueltas a Matriz de Tareas", "Transformar audios o transcripciones de reuniones en tablas estructuradas con Responsable, Tarea, Prioridad y Fecha Límite. Cero acuerdos olvidados en el equipo.", Inches(0.8)),
        ("Análisis de Requerimientos a BPMN", "Tomar la descripción verbal de un usuario de negocio y convertirla en la secuencia lógica paso a paso para diagramar y automatizar flujos en el software.", Inches(6.75))
    ]

    for title, desc, x in cards_p4:
        card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.75), Inches(5.75), Inches(5.1))
        card.fill.solid()
        card.fill.fore_color.rgb = c_card_bg
        card.line.color.rgb = c_border
        tf = card.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(20)
        p_t.font.bold = True
        p_t.font.color.rgb = c_primary
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(16)
        p_d.font.color.rgb = c_text_white
        p_d.space_before = Pt(16)

    add_speaker_note(s8, "Mencionar que el proyecto final del curso les pedirá diagramar procesos en BPMN y automatizarlos con agentes.")

    # =========================================================================
    # SLIDE 9: CHATGPT VS OPENAI PLATFORM
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_bg(s9)
    add_header(s9, "OpenAI Platform vs ChatGPT: ¿Cuál es la diferencia?")

    comp_items = [
        ("Aspecto", "ChatGPT (Consumidor)", "OpenAI Platform (API & Playground)"),
        ("Propósito", "Chatbot conversacional interactivo.", "Motor de desarrollo para automatizar sistemas."),
        ("Configuración", "Instrucciones generales.", "Control estricto de System Prompt, Temp y Tokens."),
        ("Automatización", "No se conecta a flujos externos.", "Diseñado para conectarse con Make, Zapier y n8n."),
        ("Costo", "Suscripción fija ($20/mes).", "Pago por consumo de tokens (fracciones de centavo)."),
        ("Salida de Datos", "Texto conversacional.", "Capacidad de forzar modo JSON estructurado puro.")
    ]

    rows = len(comp_items)
    cols = 3
    table_shape = s9.shapes.add_table(rows, cols, Inches(0.8), Inches(1.75), Inches(11.733), Inches(5.1))
    table = table_shape.table
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(4.766)
    table.columns[2].width = Inches(4.766)

    for r_idx, row in enumerate(comp_items):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = val
            cell.fill.solid()
            if r_idx == 0:
                cell.fill.fore_color.rgb = c_primary
                p = cell.text_frame.paragraphs[0]
                p.font.bold = True
                p.font.size = Pt(15)
                p.font.color.rgb = c_text_white
            else:
                cell.fill.fore_color.rgb = c_card_bg if r_idx % 2 == 0 else RGBColor(20, 30, 48)
                p = cell.text_frame.paragraphs[0]
                p.font.size = Pt(13.5)
                p.font.color.rgb = c_text_white if c_idx > 0 else c_accent

    add_speaker_note(s9, "Explicar que un profesional de automatización no usa la interfaz de ChatGPT para flujos de trabajo en empresas, sino la Platform y la API.")

    # =========================================================================
    # SLIDE 10: ANATOMÍA DE OPENAI PLAYGROUND
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_bg(s10)
    add_header(s10, "Los Controles de Mando en OpenAI Playground")

    controls = [
        ("System Instructions", "El rol, políticas, restricciones y formato de salida. Guía cada respuesta del agente."),
        ("User Prompt", "La consulta o texto desestructurado que ingresa el usuario o el sistema en tiempo real."),
        ("Model", "gpt-4o-mini (rápido, económico y potente para el 90% de tareas) vs gpt-4o (multimodalidad compleja)."),
        ("Temperature (0.0 a 2.0)", "0.0 = Determinista, exacto para JSON. 0.7 = Creativo para redacción y contenidos."),
        ("Max Tokens", "Límite superior de tokens en la respuesta (1 token ≈ 4 caracteres). Controla costos y extensión.")
    ]

    for idx, (t, d) in enumerate(controls):
        y = Inches(1.75 + idx * 1.02)
        card = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.7), Inches(0.92))
        card.fill.solid()
        card.fill.fore_color.rgb = c_card_bg
        card.line.color.rgb = c_border
        tf = card.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = t
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = c_accent
        p_d = tf.add_paragraph()
        p_d.text = d
        p_d.font.size = Pt(14)
        p_d.font.color.rgb = c_text_white
        p_d.space_before = Pt(2)

    add_speaker_note(s10, "Pedir a los alumnos que abran la pestaña de Playground y ubiquen cada uno de estos 5 controles.")

    # =========================================================================
    # SLIDE 11: TALLER PRÁCTICO 1 Y 2
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    add_bg(s11)
    add_header(s11, "Taller en Vivo - Prácticas 1 y 2 en Playground")

    p1_box = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.75), Inches(5.75), Inches(5.1))
    p1_box.fill.solid()
    p1_box.fill.fore_color.rgb = c_card_bg
    p1_box.line.color.rgb = c_primary
    tf_p1 = p1_box.text_frame
    tf_p1.word_wrap = True
    p = tf_p1.paragraphs[0]
    p.text = "Práctica 1: Soporte con Guardrails"
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = c_primary

    pts1 = [
        "• Modelo: gpt-4o-mini | Temp: 0.2",
        "• Rol: 'SoporteBot de TecnoHogar Perú'.",
        "• Reglas: 7 días de devolución, envíos 24-48h, liquidaciones sin cambio.",
        "• Caso de prueba: Cliente que exige devolución al día 15 amenazando con Indecopi.",
        "• Observar: Firmeza empática y activación del código [DERIVAR_A_HUMANO]."
    ]
    for pt in pts1:
        p = tf_p1.add_paragraph()
        p.text = pt
        p.font.size = Pt(14.5)
        p.font.color.rgb = c_text_white
        p.space_before = Pt(10)

    p2_box = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.75), Inches(1.75), Inches(5.75), Inches(5.1))
    p2_box.fill.solid()
    p2_box.fill.fore_color.rgb = c_card_bg
    p2_box.line.color.rgb = c_accent
    tf_p2 = p2_box.text_frame
    tf_p2.word_wrap = True
    p = tf_p2.paragraphs[0]
    p.text = "Práctica 2: Generador Multicanal"
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = c_accent

    pts2 = [
        "• Modelo: gpt-4o-mini | Temp: 0.7",
        "• Rol: 'Copywriter Estratégico B2B'.",
        "• Tarea: Generar versión LinkedIn (profesional + hashtags) y versión WhatsApp (emojis + link).",
        "• Caso de prueba: Lanzamiento de integración WhatsApp <-> Google Sheets.",
        "• Observar: Adaptación automática de longitud y registro lingüístico."
    ]
    for pt in pts2:
        p = tf_p2.add_paragraph()
        p.text = pt
        p.font.size = Pt(14.5)
        p.font.color.rgb = c_text_white
        p.space_before = Pt(10)

    add_speaker_note(s11, "Dar 10 minutos para que todos los alumnos copien los prompts de su guía y prueben respuestas en vivo.")

    # =========================================================================
    # SLIDE 12: TALLER PRÁCTICO 3 Y 4
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    add_bg(s12)
    add_header(s12, "Taller en Vivo - Prácticas 3 y 4 en Playground")

    p3_box = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.75), Inches(5.75), Inches(5.1))
    p3_box.fill.solid()
    p3_box.fill.fore_color.rgb = c_card_bg
    p3_box.line.color.rgb = c_accent_green
    tf_p3 = p3_box.text_frame
    tf_p3.word_wrap = True
    p = tf_p3.paragraphs[0]
    p.text = "Práctica 3: Extracción a JSON"
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = c_accent_green

    pts3 = [
        "• Modelo: gpt-4o-mini | Temp: 0.0",
        "• Rol: 'Motor NLP de Operaciones'.",
        "• Tarea: Analizar reclamo y devolver SOLO un JSON con: cliente, pedido, categoría, urgencia, sentimiento y acción.",
        "• Caso de prueba: Reclamo por cobro duplicado de S/ 450 en tarjeta.",
        "• Observar: Salida limpia, sin texto adicional, lista para Make/Zapier."
    ]
    for pt in pts3:
        p = tf_p3.add_paragraph()
        p.text = pt
        p.font.size = Pt(14.5)
        p.font.color.rgb = c_text_white
        p.space_before = Pt(10)

    p4_box = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.75), Inches(1.75), Inches(5.75), Inches(5.1))
    p4_box.fill.solid()
    p4_box.fill.fore_color.rgb = c_card_bg
    p4_box.line.color.rgb = c_primary
    tf_p4 = p4_box.text_frame
    tf_p4.word_wrap = True
    p = tf_p4.paragraphs[0]
    p.text = "Práctica 4: Minuta y Tareas"
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = c_primary

    pts4 = [
        "• Modelo: gpt-4o-mini | Temp: 0.3",
        "• Rol: 'Gestor de Proyectos Ágil'.",
        "• Tarea: Convertir notas informales de reunión en Minuta Ejecutiva con tabla de responsables y fechas límite.",
        "• Caso de prueba: Reunión desordenada de ventas y sistemas sobre CRM y Excel duplicado.",
        "• Observar: Síntesis de compromisos y matriz de acción."
    ]
    for pt in pts4:
        p = tf_p4.add_paragraph()
        p.text = pt
        p.font.size = Pt(14.5)
        p.font.color.rgb = c_text_white
        p.space_before = Pt(10)

    add_speaker_note(s12, "Esta práctica 3 es el corazón de la automatización moderna. Dedicar tiempo a revisar la estructura de claves y valores.")

    # =========================================================================
    # SLIDE 13: RETO APLICADO EN CLASE
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    add_bg(s13)
    add_header(s13, "Reto en Clase: Diseña el Agente de tu Trabajo (15 minutos)")

    r_card = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.75), Inches(11.7), Inches(5.1))
    r_card.fill.solid()
    r_card.fill.fore_color.rgb = c_card_bg
    r_card.line.color.rgb = c_accent
    tf_r = r_card.text_frame
    tf_r.word_wrap = True
    p = tf_r.paragraphs[0]
    p.text = "Paso a Paso del Reto Individual / Parejas:"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = c_accent

    steps = [
        "1. Identifica un proceso repetitivo de tu empleo o negocio (ej. responder cotizaciones, filtrar CVs, calificar prospectos).",
        "2. Redacta el System Prompt en tu Playground definiendo: Nombre, rol, 3 reglas indispensables y formato de salida.",
        "3. Elige la Temperatura adecuada (¿Requiere precisión numérica de datos o redacción creativa?).",
        "4. Prueba el agente con 2 casos: uno normal y un caso difícil con datos incompletos.",
        "5. Voluntarios: 2 participantes compartirán pantalla para recibir retroalimentación en vivo."
    ]
    for st in steps:
        p = tf_r.add_paragraph()
        p.text = st
        p.font.size = Pt(16)
        p.font.color.rgb = c_text_white
        p.space_before = Pt(14)

    add_speaker_note(s13, "Poner cronómetro de 10 minutos de trabajo autónomo y 5 minutos para que 2 alumnos presenten.")

    # =========================================================================
    # SLIDE 14: CONCLUSIONES Y PRÓXIMA SESIÓN
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    add_bg(s14)
    add_header(s14, "Conclusiones y Conexión con la Próxima Clase")

    c_box1 = s14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.75), Inches(5.75), Inches(5.1))
    c_box1.fill.solid()
    c_box1.fill.fore_color.rgb = c_card_bg
    c_box1.line.color.rgb = c_border
    tf_c1 = c_box1.text_frame
    tf_c1.word_wrap = True
    p = tf_c1.paragraphs[0]
    p.text = "¿Qué nos llevamos hoy? (Subindicador 3)"
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = c_accent_green

    c_pts = [
        "✓ La IA automatiza procesos de negocio cuando se le dan reglas, restricciones y guardrails claros.",
        "✓ La diferencia entre un juguete y un agente profesional radica en la arquitectura del System Prompt.",
        "✓ Dominio de parámetros técnicos: Temperature, Modelos y Tokens.",
        "✓ El formato JSON es el puente definitivo para conectar la IA con plataformas no-code (Make/n8n)."
    ]
    for cp in c_pts:
        p = tf_c1.add_paragraph()
        p.text = cp
        p.font.size = Pt(15)
        p.font.color.rgb = c_text_white
        p.space_before = Pt(12)

    c_box2 = s14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.75), Inches(1.75), Inches(5.75), Inches(5.1))
    c_box2.fill.solid()
    c_box2.fill.fore_color.rgb = c_card_bg
    c_box2.line.color.rgb = c_border
    tf_c2 = c_box2.text_frame
    tf_c2.word_wrap = True
    p = tf_c2.paragraphs[0]
    p.text = "Próxima Clase: Subindicador 4"
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = c_primary

    next_pts = [
        "• Tema: Diseño de prompts efectivos y lógica avanzada de automatización.",
        "• Técnicas de ingeniería de prompts: Few-Shot prompting, Chain of Thought (paso a paso), Delimitadores y Guardrails anti-alucinación.",
        "• Optimización e iteración de prompts para producción.",
        "• ¡Nos vemos el próximo martes a las 19:30!"
    ]
    for np in next_pts:
        p = tf_c2.add_paragraph()
        p.text = np
        p.font.size = Pt(15)
        p.font.color.rgb = c_text_white
        p.space_before = Pt(12)

    add_speaker_note(s14, "Cerrar puntuales a las 21:45. Felicitaciones a todos por sus pruebas en Playground.")

    output_path = "/Users/rpineda/Sources/Automatizacion/Sesion_Subindicador_3_Automatizacion_IA.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully with enlarged typography to {output_path}")

if __name__ == "__main__":
    create_deck()
